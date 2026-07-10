#!/usr/bin/env python3
"""
提取剧本段落结构分析。
从 .docx 中提取每个段落，识别类型（slug/character/dialog/action/parenthetical/其他），
输出 JSON 供 merge_blocks.py 使用。

用法: python3 extract_script.py <源文件.docx> [输出路径]
"""
import json, sys, re, os
from docx import Document
from docx.shared import Inches, Pt, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH

# 场景标题关键词
SLUG_KEYWORDS = [
    'INT.', 'EXT.', 'EST.', 'INT/EXT', 'EXT/INT',
    'INTERCUT', 'FADE IN', 'FADE OUT', 'FADE TO',
    'CUT TO', 'DISSOLVE', 'SMASH CUT', 'MATCH CUT',
    'FLASHBACK', 'MONTAGE', 'INSERT', 'TITLE',
    'OVER BLACK', 'OVER WHITE', 'BACK TO',
    'POV', 'OFF SCREEN', 'OS', 'V.O.', 'VO',
    'CONTINUOUS', 'LATER', 'MOMENTS LATER',
    'SERIES OF SHOTS', 'INTERCUT WITH',
]

# 常见角色名前缀（用于识别角色行）
CHARACTER_PREFIXES = [
    'นาย', 'นาง', 'นางสาว', 'คุณ', 'พี่', 'น้อง', 'เฮีย',
    'Mr.', 'Mrs.', 'Ms.', 'Dr.', 'Prof.',
]

# 对话指示词后缀（全大写角色名后常见）
DIALOG_INDICATORS = [
    '(CONT\'D)', '(CONTINUED)', '(V.O.)', '(O.S.)', '(OFF)',
    'ต่อ', 'พูด', 'เสียง',
]


def is_all_caps(text):
    """判断是否全大写（忽略标点和空格）"""
    letters = re.findall(r'[a-zA-Z\u0E00-\u0E7F]', text)  # 拉丁字母或泰文字母
    if not letters:
        return False
    latin = [c for c in letters if 'a' <= c <= 'z' or 'A' <= c <= 'Z']
    if not latin:
        return False
    return all(c.isupper() for c in latin)


def is_thai(text):
    """判断是否含泰文字符"""
    return bool(re.search(r'[\u0E00-\u0E7F]', text))


def detect_paragraph_type(para, para_idx, total_paras, prev_type=None):
    """
    识别段落类型。
    返回: (type, needs_translation)
    type: slug/character/dialog/action/parenthetical/format/note/meta/other
    """
    text = para.text.strip()
    if not text:
        return ('other', False)

    # 获取格式信息
    alignment = para.alignment
    is_bold = any(run.bold for run in para.runs if run.bold is not None)
    first_run = para.runs[0] if para.runs else None
    font_size = first_run.font.size if first_run and first_run.font.size else None
    indent_left = para.paragraph_format.left_indent
    indent_first = para.paragraph_format.first_line_indent

    # 检测是否为角色名：居中 + 全大写 + 短行
    is_centered = alignment == WD_ALIGN_PARAGRAPH.CENTER
    is_short = len(text) < 50
    all_caps = is_all_caps(text)

    # 检测是否为场景标题
    text_upper = text.upper().strip()
    slug_match = any(text_upper.startswith(kw) for kw in SLUG_KEYWORDS)
    # 或者以全大写场景关键词开头
    slug_match2 = any(text_upper.startswith(kw) for kw in [
        'INT ', 'EXT ', 'EST ', 'INT/', 'EXT/',
    ])

    # 检测括号提示
    has_parentheses = bool(re.match(r'^\s*\(', text)) or bool(re.match(r'^\s*\[', text))

    # 判断类型
    if slug_match or slug_match2:
        return ('slug', True)
    elif is_centered and all_caps and is_short:
        return ('character', False)
    elif is_centered and not all_caps and is_short:
        # 某些泰语剧本角色名不是全大写但居中
        return ('character', False)
    elif has_parentheses:
        return ('parenthetical', True)
    elif text_upper.startswith('NOTE') or text_upper.startswith('หมายเหตุ'):
        return ('note', True)
    elif text_upper.startswith('SFX') or text_upper.startswith('เสียง'):
        return ('format', True)
    elif text_upper.startswith('CUT TO') or text_upper.startswith('FADE'):
        return ('format', True)
    elif text_upper.startswith('MONTAGE') or text_upper.startswith('FLASHBACK'):
        return ('format', True)
    elif is_bold and all_caps and is_short:
        # 加粗全大写短行 → 可能是场景标题，如果没被前面的slug匹配到
        return ('slug', True)
    else:
        # 剩余：对话或动作
        # 一般对话行有缩进，动作行无缩进
        has_indent = indent_left is not None and indent_left > 0
        if has_indent:
            return ('dialog', True)
        else:
            return ('action', True)


def extract_format_info(para):
    """提取段落格式信息"""
    info = {}

    if para.alignment is not None:
        info['align'] = str(para.alignment)

    # 缩进
    if para.paragraph_format.left_indent:
        info['left_indent'] = para.paragraph_format.left_indent  # Emu

    if para.paragraph_format.first_line_indent:
        info['first_line_indent'] = para.paragraph_format.first_line_indent

    # 段落间距
    if para.paragraph_format.space_before:
        info['space_before'] = para.paragraph_format.space_before
    if para.paragraph_format.space_after:
        info['space_after'] = para.paragraph_format.space_after

    # 行距
    if para.paragraph_format.line_spacing:
        info['line_spacing'] = para.paragraph_format.line_spacing

    # 运行格式
    runs_info = []
    for run in para.runs:
        run_info = {'text': run.text}
        if run.bold is not None:
            run_info['bold'] = run.bold
        if run.italic is not None:
            run_info['italic'] = run.italic
        if run.font.size:
            run_info['size'] = str(run.font.size)
        if run.font.name:
            run_info['font'] = run.font.name
        if run.font.color and run.font.color.rgb:
            run_info['color'] = str(run.font.color.rgb)
        runs_info.append(run_info)

    info['runs'] = runs_info

    # 段落样式名
    if para.style and para.style.name:
        info['style'] = para.style.name

    return info


def main():
    docx_path = sys.argv[1] if len(sys.argv) > 1 else None
    out_path  = sys.argv[2] if len(sys.argv) > 2 else "/tmp/script_structure.json"

    if not docx_path or not os.path.exists(docx_path):
        print("用法: python3 extract_script.py <源文件.docx> [输出路径]", file=sys.stderr)
        sys.exit(1)

    doc = Document(docx_path)
    rows = []

    prev_type = None
    for idx, para in enumerate(doc.paragraphs):
        text = para.text
        if not text.strip():
            continue

        ptype, needs_trans = detect_paragraph_type(para, idx, len(doc.paragraphs), prev_type)
        fmt_info = extract_format_info(para)

        row = {
            'idx': idx,
            'type': ptype,
            'text': text,
            'needs_translation': needs_trans,
            'format': fmt_info,
        }
        rows.append(row)
        prev_type = ptype

    # 场景统计
    scene_count = sum(1 for r in rows if r['type'] == 'slug')
    char_count = sum(1 for r in rows if r['type'] == 'character')
    dialog_count = sum(1 for r in rows if r['type'] == 'dialog')
    trans_count = sum(1 for r in rows if r['needs_translation'])

    with open(out_path, 'w', encoding='utf-8') as f:
        json.dump(rows, f, ensure_ascii=False, indent=1)

    print(f"✓ 提取完成: {len(rows)} 段落")
    print(f"  ├─ 场景标题: {scene_count}")
    print(f"  ├─ 角色名:   {char_count}")
    print(f"  ├─ 对话:     {dialog_count}")
    print(f"  └─ 需翻译:   {trans_count}")
    print(f"  输出: {out_path}")


if __name__ == '__main__':
    main()